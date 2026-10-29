"""Tổng hợp báo cáo và xuất Excel. Phần 'điểm đáng chú ý' hiện sinh bằng quy tắc cố định;
sau này AI Agent sẽ viết bản nháp dựa trên cùng bộ số liệu này."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import ScoreEntry
from app.services.period import Period
from app.services.ranking_service import criterion_breakdown, get_ranking


def build_report(db: Session, period: Period) -> dict:
    ranking = get_ranking(db, period, with_change=True)
    breakdown = criterion_breakdown(db, period)
    total_entries = sum(r.entries_count for r in ranking)
    has_data = total_entries > 0

    highlights: list[str] = []
    if has_data:
        top = [r.class_name for r in ranking if r.rank == 1]
        highlights.append(f"Dẫn đầu: {', '.join(top)} với {ranking[0].total:g} điểm.")
        penalties = sorted((b for b in breakdown if b.negative < 0), key=lambda b: b.negative)
        if penalties:
            worst = penalties[0]
            highlights.append(f"Tiêu chí bị trừ nhiều nhất: {worst.name} ({abs(worst.negative):g} điểm, {worst.count} lượt chấm).")
        movers = [r for r in ranking if r.change]
        if movers:
            up = max(movers, key=lambda r: r.change)
            down = min(movers, key=lambda r: r.change)
            if up.change > 0:
                highlights.append(f"Tăng hạng nhiều nhất: {up.class_name} (tăng {up.change} hạng).")
            if down.change < 0:
                highlights.append(f"Giảm hạng nhiều nhất: {down.class_name} (giảm {abs(down.change)} hạng).")
    return {
        "period": period,
        "ranking": ranking,
        "breakdown": breakdown,
        "total_entries": total_entries,
        "total_positive": round(sum(r.positive for r in ranking), 2),
        "total_negative": round(sum(r.negative for r in ranking), 2),
        "highlights": highlights,
        "has_data": has_data,
    }


def export_excel(db: Session, period: Period) -> bytes:
    report = build_report(db, period)
    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="16213E")
    header_font = Font(bold=True, color="FFFFFF")

    def write_sheet(ws, headers, rows, widths):
        ws.append(headers)
        for cell in ws[1]:
            cell.fill, cell.font = header_fill, header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in rows:
            ws.append(row)
        for i, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "Xếp hạng"
    write_sheet(ws, ["Hạng", "Lớp", "GVCN", "Điểm cộng", "Điểm trừ", "Tổng điểm"],
                [[r.rank, r.class_name, r.homeroom_teacher, r.positive, r.negative, r.total] for r in report["ranking"]],
                [8, 14, 26, 12, 12, 12])
    ws.insert_rows(1)
    ws["A1"] = f"Bảng xếp hạng thi đua — {period.label}"
    ws["A1"].font = Font(bold=True, size=13)
    ws.freeze_panes = "A3"

    write_sheet(wb.create_sheet("Theo tiêu chí"),
                ["Tiêu chí", "Điểm cộng", "Điểm trừ", "Điểm ròng", "Số lượt chấm"],
                [[b.name, b.positive, b.negative, b.net, b.count] for b in report["breakdown"]],
                [26, 12, 12, 12, 14])

    entries = db.scalars(
        select(ScoreEntry).where(ScoreEntry.entry_date.between(period.start, period.end))
        .order_by(ScoreEntry.entry_date, ScoreEntry.id)
    ).unique().all()
    write_sheet(wb.create_sheet("Chi tiết điểm"),
                ["Ngày", "Lớp", "Tiêu chí", "Điểm", "Lý do", "Người nhập", "Thời gian nhập"],
                [[e.entry_date.strftime("%d/%m/%Y"), e.school_class.name, e.criterion.name, e.points, e.reason,
                  e.author.full_name, e.created_at.strftime("%H:%M %d/%m/%Y")] for e in entries],
                [12, 10, 18, 8, 50, 22, 18])

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
